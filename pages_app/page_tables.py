"""Page « Historique tables rafraîchies » : les tableaux du Catalogue enregistrés au moment de « Rafraîchir la table »,
la plus récente d'abord, consultables, rechargeables dans le Catalogue, téléchargeables en Excel et supprimables."""
import io

import pandas as pd
import streamlit as st

import catalogue
from titres import titre_page

titre_page("Historique tables rafraîchies", "tables")
st.caption("Chaque fois que vous cliquez sur « 🧹 Rafraîchir la table » dans le Catalogue aller-retour, le tableau "
           "affiché est enregistré ici tel qu'il était, avant d'être vidé.")

if msg := st.session_state.pop("table_supprimee", None):
    st.success(msg)
tables = catalogue.tables_rafraichies()
if not tables:
    st.info("Aucune table rafraîchie pour l'instant : utilisez « 🧹 Rafraîchir la table » sous le tableau du "
            "Catalogue aller-retour.")
    st.stop()


libelle = catalogue.libelle_table
choix = st.selectbox("Table rafraîchie", tables, format_func=libelle, key="table_rafraichie_choisie")
table = catalogue.lire_table_rafraichie(choix["nom"])


@st.cache_data(show_spinner=False)
def excel(nom: str) -> bytes:
    """La table en Excel stylé (en-tête bleu, lignes alternées, filtres), comme les autres exports."""
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
    tab = catalogue.lire_table_rafraichie(nom)
    sortie = io.BytesIO()
    with pd.ExcelWriter(sortie, engine="openpyxl") as w:
        tab.to_excel(w, sheet_name="Table rafraîchie", index=False)
        ws = w.sheets["Table rafraîchie"]
        if len(tab):
            ref = f"A1:{get_column_letter(len(tab.columns))}{len(tab) + 1}"
            t = Table(displayName="TableRafraichie", ref=ref)
            t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(t)
        for i, col in enumerate(tab.columns, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(12, len(str(col)) + 4)
    return sortie.getvalue()


st.dataframe(table, hide_index=True, width="stretch")
c0, c1, c2 = st.columns(3)
if c0.button("📂 Recharger dans le Catalogue", type="primary", width="stretch",
             disabled=catalogue.actualisation_active(),
             help="Remet cette table dans le Catalogue aller-retour pour la retravailler, avec ses champs, comme avant "
                  "« Rafraîchir la table » ; le Comparatif, la Data Visualisation et le Rapport l'utilisent de "
                  "nouveau. Si certaines de ses dates ont été recherchées depuis, leurs prix les plus récents "
                  "s'affichent."):
    catalogue.recharger_table_rafraichie(choix["nom"])
    st.session_state["table_rechargee"] = libelle(choix)
    st.switch_page("pages_app/page_catalogue.py")
c1.download_button("⬇️ Télécharger en Excel", excel(choix["nom"]), file_name=f"table_rafraichie_{choix['nom']}.xlsx",
                   mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", width="stretch")
with c2.popover("🗑️ Supprimer cette table", width="stretch"):
    st.write(f"Supprimer définitivement la table du {libelle(choix)} ?")
    if st.button("Oui, supprimer", type="primary", key=f"supprimer_{choix['nom']}"):
        catalogue.supprimer_table_rafraichie(choix["nom"])
        st.session_state["table_supprimee"] = "Table supprimée de l'historique."
        st.session_state.pop("table_rafraichie_choisie", None)
        st.rerun()
